from app.services.ingestion.pdf_loader import LoadedPage

try:
    import openpyxl
except ImportError:
    openpyxl = None


class ExcelLoader:
    """Extracts text sheet-by-sheet from Excel (.xlsx) workbooks."""

    def load(self, file_path: str) -> list[LoadedPage]:
        if openpyxl is None:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                return [LoadedPage(page_number=1, text=f.read())]

        wb = openpyxl.load_workbook(file_path, data_only=True)
        pages: list[LoadedPage] = []

        for idx, sheet_name in enumerate(wb.sheetnames, start=1):
            sheet = wb[sheet_name]
            lines = [f"Sheet: {sheet_name}"]
            for row in sheet.iter_rows(values_only=True):
                row_vals = [str(v) for v in row if v is not None]
                if row_vals:
                    lines.append(" | ".join(row_vals))

            text = "\n".join(lines)
            pages.append(LoadedPage(page_number=idx, text=text))

        wb.close()
        return pages
